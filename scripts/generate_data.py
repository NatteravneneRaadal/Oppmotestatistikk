#!/usr/bin/env python3
"""
Konverterer source-data/factStatistikk.xlsx (arket "Statistikk") til data.js
for Natteravnene Rådal-dashbordet.

Kjøres automatisk av GitHub Actions (.github/workflows/update-data.yml)
hver gang factStatistikk.xlsx endres i repositoryet. Kan også kjøres
manuelt lokalt:

    pip install openpyxl
    python3 scripts/generate_data.py
"""
import json
import datetime
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("Mangler openpyxl. Installer med: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

ROOT = Path(__file__).resolve().parent.parent
EXCEL_PATH = ROOT / "source-data" / "factStatistikk.xlsx"
OUTPUT_PATH = ROOT / "data.js"

REAL_SCHOOLS = {
    'Rå', 'Rådalslien', 'Skranevatnet', 'Skeie', 'Skjold',
    'Søråshøgda', 'Aurdalslia', 'Søreide', 'Apeltun'
}
SCHOOL_ORDER = [
    'Rå', 'Rådalslien', 'Skranevatnet', 'Skeie', 'Skjold',
    'Søråshøgda', 'Aurdalslia', 'Søreide', 'Apeltun'
]
SKOLE_IDX = {s: i for i, s in enumerate(SCHOOL_ORDER)}
KATEGORI_LISTE = ['Foreldre', 'Frivillige', 'MC ravn']


def main():
    if not EXCEL_PATH.exists():
        print(f"Fant ikke Excel-filen: {EXCEL_PATH}", file=sys.stderr)
        sys.exit(1)

    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    if "Statistikk" not in wb.sheetnames:
        print('Fant ikke arket "Statistikk" i Excel-filen.', file=sys.stderr)
        sys.exit(1)
    ws = wb["Statistikk"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))

    records = []
    skipped = 0
    for r in rows:
        if len(r) < 7:
            continue
        dato, skole, innkalt, mott, aar, semester, sesong = r[:7]
        if dato is None or skole is None or sesong is None:
            continue
        # "Kl 20" / "Kl 22" er et alternativt tidspunkt-oppsett (kun høsten 2024)
        # som dobbelttegner de samme vaktene som allerede telles pr skole.
        if skole in ('Kl 20', 'Kl 22'):
            skipped += 1
            continue
        if skole in REAL_SCHOOLS:
            kat_i = 0
            sk_i = SKOLE_IDX[skole]
        elif skole == 'Frivillige':
            kat_i = 1
            sk_i = -1
        elif skole == 'MC ravn':
            kat_i = 2
            sk_i = -1
        else:
            skipped += 1
            continue

        if not hasattr(dato, 'strftime'):
            continue

        records.append((dato, sk_i, kat_i, innkalt, mott, sesong))

    if not records:
        print("Ingen gyldige rader funnet i Excel-filen.", file=sys.stderr)
        sys.exit(1)

    seasons = sorted(
        set(r[5] for r in records),
        key=lambda s: (int(s.split()[1]), 0 if s.startswith('Vår') else 1)
    )
    sesong_idx = {s: i for i, s in enumerate(seasons)}

    compact_rows = []
    for dato, sk_i, kat_i, innkalt, mott, sesong in records:
        datoint = int(dato.strftime('%Y%m%d'))
        compact_rows.append([datoint, sk_i, kat_i, innkalt, mott, sesong_idx[sesong]])

    out = {
        'generert': datetime.date.today().strftime('%Y-%m-%d'),
        'skoleRekkefolge': SCHOOL_ORDER,
        'sesongRekkefolge': seasons,
        'kategoriListe': KATEGORI_LISTE,
        'rows': compact_rows
    }

    compact = json.dumps(out, ensure_ascii=False, separators=(',', ':'))
    # Formater 'rows' over flere kortere linjer (i stedet for én svært lang linje)
    # - dette er lettere for ulike verktøy/forhåndsvisninger å håndtere.
    rows_json = out['rows']
    chunk_size = 15
    chunk_lines = []
    for i in range(0, len(rows_json), chunk_size):
        chunk = rows_json[i:i + chunk_size]
        chunk_lines.append(','.join(json.dumps(r, separators=(',', ':')) for r in chunk))
    rows_str = '[\n' + ',\n'.join(chunk_lines) + '\n]'

    meta = {k: v for k, v in out.items() if k != 'rows'}
    meta_str = json.dumps(meta, ensure_ascii=False, separators=(',', ':'))
    compact = meta_str[:-1] + ',"rows":' + rows_str + '}'

    js_content = (
        "// Datakilde for Natteravnene Rådal - statistikk\n"
        "// Denne filen genereres automatisk fra source-data/factStatistikk.xlsx\n"
        "// av GitHub Actions - ikke rediger den manuelt.\n"
        "const RAVNEDATA = " + compact + ";\n"
    )
    OUTPUT_PATH.write_text(js_content, encoding='utf-8')

    print(f"OK: {len(records)} rader skrevet til {OUTPUT_PATH} ({skipped} rader hoppet over).")


if __name__ == "__main__":
    main()
